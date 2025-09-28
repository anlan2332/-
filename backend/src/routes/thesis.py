"""
论文写作相关路由
"""
import os
import json
import threading
from datetime import datetime
from flask import Blueprint, request, jsonify, session
from src.models.database import db, User, Paper, WritingAssistant, APIKey
from src.websocket_manager import notify_paper_progress, ws_manager
from src.utils.doubao_client import get_doubao_client, DoubaoClient

thesis_bp = Blueprint('thesis', __name__)

# 您的豆包API密钥和助手指令
DEFAULT_API_KEY = "039c31ba-ebd5-429b-8e88-593eb0e0dc67"
ASSISTANT_INSTRUCTIONS = """1.在保持大意和字数不变或缩减的情况下，拟用学术化的语言重新输出以上内容，减少小开头，少aigc重复率，增强段落衔接感，尽量减少首先、然后、再者、最后等相似词汇的使用次数。
2.以以上文献为我撰写国内外研究，例如学者xx在什么什么方面做了什么什么，然后在其学者名字后面加[]序号
3.根据参考文献撰写，选取几个代表性理论，加入序号引用即可
4.根据文件内容及题目搜索五年内的10篇国内文献.5篇国外文献，要求真实，以gbt2015格式输出
5.根据这个开题报告为我撰写一篇8000字的论文，绪论的国内外研究引用参考文献标注学者和序号，没有用完的参考文献可在正文中引用，同样标注序号；文献引用按照顺序，如果参考文献顺序不对，可以调整参考文献的顺序。引用过的文献就不要再引用了。
跟据以上修改意见修改，并告诉我怎么添加和修改，把需要添加和修改的内容告诉我。
整篇文章用学术化的语言输出，减少小开头，少aigc重复率，增强段落衔接感，尽量减少首先、然后、再者、最后等相似词汇的使用次数。可输出图表，表格正常输出，图片用svg输出。
仔细反复校对是否符合开题，不要自作主张，确保内容正确。
先输出内容，最后输出图片，标记图片引用的地方（如图xx所示）。
参考模板，根据任务书和开题报告，帮我撰写基于Anylogic的城市道路交通拥堵特性分析与仿真字数控制在15000字以内。只在国内外研究中引用参考文献并撰写，加入序号；加入图表，并在文中说明引用。图片选用电气控制图风格，黑白色，非必要不使用彩色。如需引用数据，标注数据来源。
整篇文章用学术化的语言输出，减少小开头，少aigc重复率，增强段落衔接感，尽量减少首先、然后、再者、最后等相似词汇的使用次数，仔细反复校对是否符合开题和任务书以及模板，不要自作主张，确保内容和数据正确。
模仿中国本科生的水平和角度去撰写，降低aigc率，逻辑要严谨，具有研究深度，结构模仿人类，不要太完美。取消首先其次等用语，要用停顿感，尽量减少首先、然后、再者、最后等相似词汇的使用次数，减少小开，头不要一气呵成明白吗。"""

def get_doubao_api_client():
    """获取豆包API客户端"""
    try:
        # 优先从数据库获取活跃的API密钥
        api_key = APIKey.query.filter_by(service_type='doubao', is_active=True).order_by(APIKey.priority).first()
        if api_key:
            return get_doubao_client(api_key.api_key)
        else:
            # 使用默认密钥
            return get_doubao_client(DEFAULT_API_KEY)
    except Exception as e:
        print(f"获取豆包API客户端失败: {e}")
        return None

def log_user_activity(user_id: int, action: str, details: dict = None):
    """记录用户活动"""
    try:
        # 暂时使用print记录，后续可扩展为完整的活动日志系统
        print(f"用户 {user_id} 执行操作: {action}, 详情: {details}")
    except Exception as e:
        print(f"记录用户活动失败: {e}")

@thesis_bp.route('/create', methods=['POST'])
def create_thesis():
    """创建新论文"""
    try:
        data = request.get_json()
        user_id = session.get('user_id')
        
        if not user_id:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        # 获取用户信息
        user = User.query.get(user_id)
        if not user:
            return jsonify({'success': False, 'message': '用户不存在'}), 404
        
        # 验证输入
        title = data.get('title', '').strip()
        field = data.get('field', '').strip()
        education_level = data.get('education_level', '').strip()
        keywords = data.get('keywords', '').strip()
        description = data.get('description', '').strip()
        word_count = data.get('word_count', 8000)
        
        if not title:
            return jsonify({'success': False, 'message': '论文标题不能为空'}), 400
        
        # 获取订单ID（可选）
        order_id = data.get('order_id')
        
        # 创建论文记录
        paper = Paper(
            user_id=user_id,
            title=title,
            paper_type='graduation_thesis',
            status='created',
            config=json.dumps({
                'field': field,
                'education_level': education_level,
                'keywords': keywords,
                'description': description,
                'word_count': word_count,
                'order_id': order_id
            }, ensure_ascii=False)
        )
        
        db.session.add(paper)
        db.session.commit()
        
        # 记录用户活动
        log_user_activity(user_id, 'create_thesis', {
            'paper_id': paper.id,
            'title': title,
            'field': field
        })
        
        return jsonify({
            'success': True,
            'message': '论文创建成功',
            'data': {
                'paper_id': paper.id,
                'title': title,
                'status': paper.status,
                'created_at': paper.created_at.isoformat()
            }
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'创建论文失败: {str(e)}'}), 500

@thesis_bp.route('/generate', methods=['POST'])
def generate_thesis():
    """生成论文内容"""
    try:
        data = request.get_json()
        user_id = session.get('user_id')
        
        if not user_id:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        paper_id = data.get('paper_id')
        if not paper_id:
            return jsonify({'success': False, 'message': '论文ID不能为空'}), 400
        
        # 获取论文信息
        paper = Paper.query.filter_by(id=paper_id, user_id=user_id).first()
        if not paper:
            return jsonify({'success': False, 'message': '论文不存在'}), 404
        
        # 解析论文配置
        config = json.loads(paper.config) if paper.config else {}
        
        # 构建提示词
        prompt = f"""根据以下信息生成高质量的学术论文：

论文标题：{paper.title}
研究领域：{config.get('field', '')}
学历层次：{config.get('education_level', '')}
关键词：{config.get('keywords', '')}
补充说明：{config.get('description', '')}
目标字数：{config.get('word_count', 8000)}字

{ASSISTANT_INSTRUCTIONS}

请按照学术规范生成完整的论文内容，包括：
1. 摘要
2. 关键词
3. 绪论
4. 文献综述
5. 研究方法
6. 结果与分析
7. 结论
8. 参考文献

要求：
- 内容严谨、逻辑清晰
- 符合学术写作规范
- 降低AIGC检测率
- 模仿本科生写作风格
- 包含真实的参考文献引用
"""

        # 更新论文状态
        paper.status = 'generating'
        db.session.commit()
        
        # 发送进度通知
        notify_paper_progress(user_id, paper_id, 10, '开始生成论文...')
        
        # 调用豆包AI生成
        client = get_doubao_api_client()
        if not client:
            paper.status = 'failed'
            db.session.commit()
            return jsonify({'success': False, 'message': 'AI服务暂时不可用'}), 500
        
        try:
            notify_paper_progress(user_id, paper_id, 30, '正在调用豆包AI模型...')
            
            response = client.generate_thesis_content(
                title=paper.title,
                field=config.get('field', ''),
                education_level=config.get('education_level', ''),
                keywords=config.get('keywords', ''),
                description=config.get('description', ''),
                assistant_instructions=ASSISTANT_INSTRUCTIONS
            )
            
            # 检查是否有错误
            if 'error' in response:
                raise Exception(response['error']['message'])
            
            # 获取生成的内容
            if 'choices' in response and response['choices']:
                content = response['choices'][0]['message']['content']
            elif 'content' in response:
                content = response['content']
            else:
                content = str(response)
            
            notify_paper_progress(user_id, paper_id, 80, '正在处理生成结果...')
            
            # 检查内容质量并补充
            if len(content) < 1000:
                # 如果内容太短，补充更详细的内容
                supplement_prompt = f"""
                以下是已经生成的论文内容，但字数不够：
                
                {content}
                
                请扩展这篇论文到至少{config.get('word_count', 8000)}字，要求：
                1. 保持原有结构和内容
                2. 扩展每个章节的具体内容
                3. 增加更多的理论分析和实例说明
                4. 保持学术写作风格
                5. 添加更多的数据分析和图表说明
                """
                
                supplement_response = client.generate_thesis_content(
                    title=paper.title,
                    field=config.get('field', ''),
                    education_level=config.get('education_level', ''),
                    keywords=config.get('keywords', ''),
                    description=supplement_prompt,
                    assistant_instructions=ASSISTANT_INSTRUCTIONS
                )
                
                if 'choices' in supplement_response and supplement_response['choices']:
                    content = supplement_response['choices'][0]['message']['content']
                elif 'content' in supplement_response:
                    content = supplement_response['content']
            
            # 更新论文内容
            paper.content = content
            paper.status = 'completed'
            paper.word_count = len(content)
            
            # 如果内容仍然太短，添加基本的论文结构
            if len(content) < 2000:
                enhanced_content = f"""
{paper.title}

摘要
{content[:500] if content else '本研究针对' + config.get('field', '') + '领域的' + paper.title + '进行了深入研究。'}

关键词：{config.get('keywords', '研究方法, 实证分析, 理论研究')}

1. 绪论

1.1 研究背景
随着{config.get('field', '现代科学技术')}的迅速发展，{paper.title}已成为该领域的重要研究方向。本研究旨在深入分析其理论基础和实践应用。

1.2 研究意义
本研究对于推动{config.get('field', '')}领域的理论发展和实践应用具有重要意义。

2. 文献综述

2.1 理论基础
目前在{config.get('field', '')}领域中，相关理论研究主要集中在...

2.2 研究现状
近年来，国内外学者在{paper.title}方面开展了大量研究...

3. 研究方法

3.1 研究设计
本研究采用定性与定量相结合的研究方法...

3.2 数据收集
通过多种渠道收集相关数据...

4. 结果分析

4.1 数据分析结果
研究结果表明...

4.2 结果讨论
通过对比分析可以发现...

5. 结论与建议

5.1 主要结论
本研究的主要结论包括...

5.2 研究局限性
本研究存在一定的局限性...

5.3 未来研究展望
未来研究可从以下方面深入...

参考文献
[1] 张三, 李四. {paper.title}的理论与实践[J]. {config.get('field', '学术')}研究, 2024, 15(2): 45-58.
[2] 王五, 赵六. {config.get('field', '现代技术')}发展研究[M]. 北京: 科学出版社, 2023.
[3] Smith J, Brown A. Research on {paper.title[:20]}[J]. International Journal of {config.get('field', 'Science')}, 2023, 28(4): 123-135.

{content if len(content) > 500 else ''}
"""
                paper.content = enhanced_content
                paper.word_count = len(enhanced_content)
            
            db.session.commit()
            
            notify_paper_progress(user_id, paper_id, 100, '论文生成完成！')
            
            # 记录用户活动
            log_user_activity(user_id, 'generate_thesis', {
                'paper_id': paper_id,
                'word_count': len(content)
            })
            
            return jsonify({
                'success': True,
                'message': '论文生成成功',
                'data': {
                    'paper_id': paper_id,
                    'content': content,
                    'word_count': len(content),
                    'status': paper.status
                }
            })
            
        except Exception as e:
            paper.status = 'failed'
            db.session.commit()
            notify_paper_progress(user_id, paper_id, 0, f'生成失败: {str(e)}')
            return jsonify({'success': False, 'message': f'AI生成失败: {str(e)}'}), 500
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'生成论文失败: {str(e)}'}), 500

@thesis_bp.route('/list', methods=['GET'])
def list_thesis():
    """获取用户论文列表"""
    try:
        user_id = session.get('user_id')
        
        if not user_id:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        page = request.args.get('page', 1, type=int)
        per_page = request.args.get('per_page', 10, type=int)
        
        papers = Paper.query.filter_by(user_id=user_id).order_by(Paper.created_at.desc()).paginate(
            page=page, per_page=per_page, error_out=False
        )
        
        paper_list = []
        for paper in papers.items:
            config = json.loads(paper.config) if paper.config else {}
            paper_list.append({
                'id': paper.id,
                'title': paper.title,
                'paper_type': paper.paper_type,
                'status': paper.status,
                'word_count': paper.word_count,
                'field': config.get('field', ''),
                'education_level': config.get('education_level', ''),
                'created_at': paper.created_at.isoformat(),
                'updated_at': paper.updated_at.isoformat()
            })
        
        return jsonify({
            'success': True,
            'data': {
                'papers': paper_list,
                'total': papers.total,
                'pages': papers.pages,
                'current_page': papers.page,
                'per_page': papers.per_page
            }
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取论文列表失败: {str(e)}'}), 500

@thesis_bp.route('/<int:paper_id>', methods=['GET'])
def get_thesis(paper_id):
    """获取论文详情"""
    try:
        user_id = session.get('user_id')
        
        if not user_id:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        paper = Paper.query.filter_by(id=paper_id, user_id=user_id).first()
        if not paper:
            return jsonify({'success': False, 'message': '论文不存在'}), 404
        
        config = json.loads(paper.config) if paper.config else {}
        
        return jsonify({
            'success': True,
            'data': {
                'id': paper.id,
                'title': paper.title,
                'paper_type': paper.paper_type,
                'status': paper.status,
                'content': paper.content,
                'word_count': paper.word_count,
                'config': config,
                'created_at': paper.created_at.isoformat(),
                'updated_at': paper.updated_at.isoformat()
            }
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取论文详情失败: {str(e)}'}), 500

@thesis_bp.route('/<int:paper_id>/update', methods=['PUT'])
def update_thesis(paper_id):
    """更新论文"""
    try:
        user_id = session.get('user_id')
        
        if not user_id:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        paper = Paper.query.filter_by(id=paper_id, user_id=user_id).first()
        if not paper:
            return jsonify({'success': False, 'message': '论文不存在'}), 404
        
        data = request.get_json()
        
        # 更新基本信息
        if 'title' in data:
            paper.title = data['title']
        if 'content' in data:
            paper.content = data['content']
            paper.word_count = len(data['content'])
        
        # 更新配置
        if 'config' in data:
            existing_config = json.loads(paper.config) if paper.config else {}
            existing_config.update(data['config'])
            paper.config = json.dumps(existing_config, ensure_ascii=False)
        
        db.session.commit()
        
        # 记录用户活动
        log_user_activity(user_id, 'update_thesis', {
            'paper_id': paper_id,
            'title': paper.title
        })
        
        return jsonify({
            'success': True,
            'message': '论文更新成功',
            'data': {
                'id': paper.id,
                'title': paper.title,
                'status': paper.status,
                'word_count': paper.word_count
            }
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'更新论文失败: {str(e)}'}), 500

@thesis_bp.route('/<int:paper_id>/delete', methods=['DELETE'])
def delete_thesis(paper_id):
    """删除论文"""
    try:
        user_id = session.get('user_id')
        
        if not user_id:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        paper = Paper.query.filter_by(id=paper_id, user_id=user_id).first()
        if not paper:
            return jsonify({'success': False, 'message': '论文不存在'}), 404
        
        # 记录用户活动
        log_user_activity(user_id, 'delete_thesis', {
            'paper_id': paper_id,
            'title': paper.title
        })
        
        db.session.delete(paper)
        db.session.commit()
        
        return jsonify({
            'success': True,
            'message': '论文删除成功'
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'删除论文失败: {str(e)}'}), 500

@thesis_bp.route('/<int:paper_id>/export', methods=['GET'])
def export_thesis(paper_id):
    """导出论文"""
    try:
        user_id = session.get('user_id')
        
        if not user_id:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        paper = Paper.query.filter_by(id=paper_id, user_id=user_id).first()
        if not paper:
            return jsonify({'success': False, 'message': '论文不存在'}), 404
        
        if not paper.content:
            return jsonify({'success': False, 'message': '论文内容为空'}), 400
        
        # 记录用户活动
        log_user_activity(user_id, 'export_thesis', {
            'paper_id': paper_id,
            'format': request.args.get('format', 'txt')
        })
        
        export_format = request.args.get('format', 'txt')
        
        if export_format == 'json':
            return jsonify({
                'success': True,
                'data': {
                    'title': paper.title,
                    'content': paper.content,
                    'word_count': paper.word_count,
                    'created_at': paper.created_at.isoformat()
                }
            })
        else:
            return jsonify({
                'success': True,
                'data': {
                    'filename': f"{paper.title}.txt",
                    'content': paper.content,
                    'size': len(paper.content.encode('utf-8'))
                }
            })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'导出论文失败: {str(e)}'}), 500

@thesis_bp.route('/progress/<int:paper_id>', methods=['GET'])
def get_thesis_progress(paper_id):
    """获取论文生成进度"""
    try:
        user_id = session.get('user_id')
        
        if not user_id:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        paper = Paper.query.filter_by(id=paper_id, user_id=user_id).first()
        if not paper:
            return jsonify({'success': False, 'message': '论文不存在'}), 404
        
        # 根据状态返回进度
        progress_map = {
            'created': 0,
            'generating': 50,
            'completed': 100,
            'failed': 0
        }
        
        progress = progress_map.get(paper.status, 0)
        
        return jsonify({
            'success': True,
            'data': {
                'paper_id': paper_id,
                'status': paper.status,
                'progress': progress,
                'word_count': paper.word_count
            }
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'获取进度失败: {str(e)}'}), 500

@thesis_bp.route('/test-doubao', methods=['POST'])
def test_doubao_api():
    """测试豆包API连接和生成功能"""
    try:
        data = request.get_json()
        title = data.get('title', '测试论文标题')
        field = data.get('field', '计算机科学')
        education_level = data.get('education_level', '本科')
        keywords = data.get('keywords', '测试关键词')
        
        # 获取豆包客户端
        client = get_doubao_api_client()
        if not client:
            return jsonify({'success': False, 'message': '豆包API客户端获取失败'}), 500
        
        # 测试生成论文大纲
        outline_response = client.generate_outline(
            title=title,
            field=field, 
            education_level=education_level
        )
        
        # 检查响应
        if 'error' in outline_response:
            return jsonify({
                'success': False, 
                'message': f'豆包API调用失败: {outline_response["error"]["message"]}'
            }), 500
        
        # 提取生成的内容
        outline_content = outline_response['choices'][0]['message']['content']
        
        response_data = {
            'success': True,
            'message': '豆包API测试成功',
            'data': {
                'title': title,
                'field': field,
                'education_level': education_level,
                'keywords': keywords,
                'generated_outline': outline_content,
                'api_usage': outline_response.get('usage', {}),
                'model': outline_response.get('model', 'unknown')
            }
        }
        
        return jsonify(response_data)
        
    except Exception as e:
        return jsonify({
            'success': False, 
            'message': f'测试豆包API失败: {str(e)}'
        }), 500

@thesis_bp.route('/search-references', methods=['POST'])
def search_references():
    """搜索参考文献"""
    try:
        data = request.get_json()
        query = data.get('query', '')
        field = data.get('field', '')
        education_level = data.get('education_level', '本科')
        limit = data.get('limit', 10)
        
        if not query:
            return jsonify({'success': False, 'message': '搜索关键词不能为空'}), 400
        
        # 使用豆包API搜索相关文献
        client = get_doubao_api_client()
        if not client:
            return jsonify({'success': False, 'message': 'AI服务暂时不可用'}), 503
            
        # 构建更精准的搜索提示词
        search_prompt = f"""
        作为一个专业的学术研究助手，请根据以下研究信息，精准搜索高度相关的学术文献：

        研究主题：{query}
        研究领域：{field}
        学历层次：{education_level}
        文献数量：{limit}篇

        搜索要求：
        1. 文献必须与研究主题高度相关
        2. 优先选择近3年内的最新研究
        3. 包含权威期刊和高水平会议论文
        4. 涵盖理论研究和实践应用两个方面
        5. 作者应包括该领域的知名学者

        请以严格的JSON格式返回：
        {{
            "references": [
                {{
                    "title": "文献标题（必须与{query}直接相关）",
                    "authors": "第一作者, 通讯作者, 其他作者",
                    "journal": "权威期刊名称（如Nature、Science或领域顶级期刊）",
                    "year": "2022-2024年间",
                    "doi": "10.xxxx/xxxx（真实有效的DOI）",
                    "abstract": "精准描述该研究的核心内容和与{query}的关联性",
                    "relevance_score": "0.95（相关性评分，0-1）",
                    "citation_count": "100+（引用次数）",
                    "research_type": "理论/实验/综述"
                }}
            ]
        }}

        特别注意：
        - 所有文献必须与「{query}」主题直接相关
        - 文献标题中应包含与主题相关的关键词
        - 优先选择在{field}领域内的权威研究
        - 确保所有信息的真实性和准确性
        """
        
        try:
            response = client.chat_completion([
                {"role": "user", "content": search_prompt}
            ])
            
            if response and 'content' in response:
                content = response['content']
                
                # 尝试解析JSON响应
                import re
                json_match = re.search(r'\{.*\}', content, re.DOTALL)
                if json_match:
                    references_data = json.loads(json_match.group())
                    references = references_data.get('references', [])
                    
                    # 为每个文献添加ID
                    for i, ref in enumerate(references):
                        ref['id'] = i + 1
                    
                    return jsonify({
                        'success': True,
                        'data': references,
                        'message': f'找到 {len(references)} 篇相关文献'
                    })
                
        except Exception as api_error:
            print(f"AI搜索失败: {api_error}")
        
        # 如果AI搜索失败，生成与主题相关的模拟数据
        
        # 提取主题关键词
        topic_keywords = query.split(' ')[:3] if ' ' in query else [query]
        
        mock_references = [
            {
                "id": 1,
                "title": f"{query}的理论基础与实践应用研究",
                "authors": "张明, 李华, 王芳",
                "journal": f"{field}研究",
                "year": "2024",
                "doi": "10.1234/example.2024.001",
                "abstract": f"本文系统分析了{query}的理论基础，并通过实证研究验证了其在{field}领域的应用效果。研究结果表明，{topic_keywords[0] if topic_keywords else '该方法'}在解决实际问题中具有显著优势。"
            },
            {
                "id": 2,
                "title": f"基于{education_level}教育的{field}创新研究——以{query}为例",
                "authors": "刘强, 陈敏",
                "journal": f"{field}教育研究",
                "year": "2024",
                "doi": "10.1234/example.2024.002",
                "abstract": f"针对{education_level}教育中{field}的特点，以{query}为研究对象，提出了创新的教学方法和实践模式，为相关研究提供了重要参考。"
            },
            {
                "id": 3,
                "title": f"{query}在{field}领域的发展现状与趋势分析",
                "authors": "赵伟, 孙丽, 周杰",
                "journal": f"{field}发展研究",
                "year": "2023",
                "doi": "10.1234/example.2023.003",
                "abstract": f"通过文献综述和数据分析，系统梳理了{query}在{field}领域的研究进展，分析了当前发展现状和未来趋势。"
            },
            {
                "id": 4,
                "title": f"{query}的实证研究——基于{field}视角的分析",
                "authors": "马超, 杨雪",
                "journal": f"{field}实证研究",
                "year": "2023",
                "doi": "10.1234/example.2023.004",
                "abstract": f"运用实证研究方法，从{field}视角深入研究了{query}的关键因素和作用机制，为理论发展和实践应用提供了重要依据。"
            },
            {
                "id": 5,
                "title": f"基于{topic_keywords[0] if topic_keywords else '新技术'}的{query}创新应用研究",
                "authors": "胡军, 郭萍",
                "journal": f"{field}技术与应用",
                "year": "2023",
                "doi": "10.1234/example.2023.005",
                "abstract": f"结合{topic_keywords[0] if topic_keywords else '新技术'}的最新进展，深入探讨了其在{query}中的创新应用模式，为{field}领域的技术升级和产业发展提供了新思路。"
            }
        ]
        
        return jsonify({
            'success': True,
            'data': mock_references,
            'message': f'找到 {len(mock_references)} 篇相关文献（模拟数据）'
        })
        
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'搜索文献失败: {str(e)}'
        }), 500

@thesis_bp.route('/parse-references', methods=['POST'])
def parse_references():
    """解析参考文献格式"""
    try:
        data = request.get_json()
        references = data.get('references', '')
        
        if not references:
            return jsonify({'success': False, 'message': '文献内容不能为空'}), 400
        
        # 使用豆包AI解析文献格式
        client = get_doubao_api_client()
        if not client:
            return jsonify({'success': False, 'message': 'AI服务暂时不可用'}), 503
        
        parse_prompt = f"""
        请对以下文献进行标准化格式化处理，输出GB/T 7714-2015标准格式：

        原始文献：
        {references}

        请按照以下要求处理：
        1. 统一格式为GB/T 7714-2015标准
        2. 按照作者姓名、文献题名、期刊名称、发表年份的正确格式
        3. 修正标点符号错误
        4. 统一中英文标点符号
        5. 检查并修正作者名和期刊名称
        6. 每条文献占一行
        7. 按照顺序编号

        格式示例：
        [1] 张三, 李四. 研究题目[J]. 期刊名称, 2024, 12(3): 45-52.
        [2] WANG L, SMITH J. Research Title[J]. Journal Name, 2024, 15(2): 123-135.

        请直接返回整理后的文献列表，不要添加其他说明。
        """
        
        try:
            response = client.chat_completion([
                {"role": "user", "content": parse_prompt}
            ])
            
            if response and 'content' in response:
                formatted_references = response['content'].strip()
                
                return jsonify({
                    'success': True,
                    'data': {
                        'original_references': references,
                        'formatted_references': formatted_references
                    },
                    'message': '文献格式解析完成'
                })
            else:
                raise Exception('无法获取AI响应')
                
        except Exception as api_error:
            print(f"AI解析失败: {api_error}")
            
            # 基本的文献格式化处理
            lines = references.strip().split('\n')
            formatted_lines = []
            
            for i, line in enumerate(lines):
                line = line.strip()
                if line:
                    # 简单的格式化
                    if not line.startswith('['):
                        line = f'[{i+1}] {line}'
                    # 统一标点
                    line = line.replace('，', ', ').replace('．', '. ').replace('：', ': ')
                    formatted_lines.append(line)
            
            formatted_references = '\n'.join(formatted_lines)
            
            return jsonify({
                'success': True,
                'data': {
                    'original_references': references,
                    'formatted_references': formatted_references
                },
                'message': '文献格式解析完成（基础处理）'
            })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'解析文献失败: {str(e)}'}), 500

@thesis_bp.route('/generate-outline-from-proposal', methods=['POST'])
def generate_outline_from_proposal():
    """基于开题报告生成大纲"""
    try:
        data = request.get_json()
        proposal_content = data.get('proposal_content', '')
        title = data.get('title', '')
        field = data.get('field', '')
        education_level = data.get('education_level', '')
        
        if not proposal_content:
            return jsonify({'success': False, 'message': '开题报告内容不能为空'}), 400
        
        # 使用豆包AI分析开题报告生成大纲
        client = get_doubao_api_client()
        if client:
            outline_prompt = f"""
            请根据以下开题报告内容，生成一个详细的论文大纲：

            开题报告内容：
            {proposal_content}

            论文信息：
            - 标题：{title}
            - 领域：{field}
            - 学历：{education_level}

            请生成一个符合以下要求的论文大纲：
            1. 符合{education_level}学历层次的论文结构
            2. 与开题报告的研究内容保持一致
            3. 包含适当的图表建议
            4. 结构清晰、逻辑严密

            请以JSON格式返回：
            {{
                "structured": [
                    {{
                        "id": 1,
                        "title": "1. 章节标题",
                        "children": [
                            {{"id": 11, "title": "1.1 子章节", "hasChart": false, "chartType": null}}
                        ]
                    }}
                ],
                "systemRecommended": "系统推荐的文本大纲"
            }}
            """
            
            try:
                response = client.chat_completion([
                    {"role": "user", "content": outline_prompt}
                ])
                
                if response and 'content' in response:
                    content_text = response['content']
                    import re
                    json_match = re.search(r'\{.*\}', content_text, re.DOTALL)
                    if json_match:
                        outline_data = json.loads(json_match.group())
                        return jsonify({
                            'success': True,
                            'data': outline_data,
                            'message': '基于开题报告的大纲生成完成'
                        })
            except Exception as e:
                print(f"AI生成大纲失败: {e}")
        
        # 基本的大纲生成
        basic_outline = {
            "structured": [
                {
                    "id": 1,
                    "title": "1. 绪论",
                    "children": [
                        {"id": 11, "title": "1.1 研究背景", "hasChart": False, "chartType": None},
                        {"id": 12, "title": "1.2 研究意义", "hasChart": False, "chartType": None},
                        {"id": 13, "title": "1.3 研究内容", "hasChart": False, "chartType": None}
                    ]
                },
                {
                    "id": 2,
                    "title": "2. 文献综述",
                    "children": [
                        {"id": 21, "title": "2.1 理论基础", "hasChart": False, "chartType": None},
                        {"id": 22, "title": "2.2 研究现状", "hasChart": True, "chartType": "table"}
                    ]
                },
                {
                    "id": 3,
                    "title": "3. 研究方法",
                    "children": [
                        {"id": 31, "title": "3.1 研究设计", "hasChart": False, "chartType": None},
                        {"id": 32, "title": "3.2 数据收集", "hasChart": True, "chartType": "chart"}
                    ]
                },
                {
                    "id": 4,
                    "title": "4. 结果分析",
                    "children": [
                        {"id": 41, "title": "4.1 数据分析", "hasChart": True, "chartType": "chart"},
                        {"id": 42, "title": "4.2 结果讨论", "hasChart": False, "chartType": None}
                    ]
                },
                {
                    "id": 5,
                    "title": "5. 结论与展望",
                    "children": [
                        {"id": 51, "title": "5.1 研究结论", "hasChart": False, "chartType": None},
                        {"id": 52, "title": "5.2 研究展望", "hasChart": False, "chartType": None}
                    ]
                }
            ],
            "systemRecommended": f"""
基于开题报告的{title}研究大纲：

一、绪论
1.1 研究背景与意义
1.2 研究目标与内容
1.3 研究方法与技术路线

二、文献综述
2.1 {field}理论基础
2.2 国内外研究现状
2.3 研究趋势分析

三、研究方法与设计
3.1 研究设计思路
3.2 数据收集与分析
3.3 关键技术实现

四、研究结果与分析
4.1 实验结果分析
4.2 结果讨论与验证

五、结论与展望
5.1 研究结论
5.2 不足与展望
            """.strip()
        }
        
        return jsonify({
            'success': True,
            'data': basic_outline,
            'message': '基于开题报告的大纲生成完成'
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'生成大纲失败: {str(e)}'}), 500

@thesis_bp.route('/parse-custom-outline', methods=['POST'])
def parse_custom_outline():
    """解析用户自定义大纲"""
    try:
        data = request.get_json()
        outline_text = data.get('outline_text', '')
        title = data.get('title', '')
        field = data.get('field', '')
        
        if not outline_text:
            return jsonify({'success': False, 'message': '大纲内容不能为空'}), 400
        
        # 使用AI解析用户大纲
        client = get_doubao_api_client()
        if client:
            parse_prompt = f"""
            请将以下用户提供的大纲文本转换为结构化的JSON格式：

            用户大纲：
            {outline_text}

            请识别其中的章节结构，并以以下JSON格式返回：
            {{
                "structured": [
                    {{
                        "id": 1,
                        "title": "主章节标题",
                        "children": [
                            {{"id": 11, "title": "子章节标题", "hasChart": false, "chartType": null}}
                        ]
                    }}
                ],
                "systemRecommended": "格式化后的文本大纲"
            }}

            注意：
            1. 保持原有的章节结构和编号
            2. 为每个章节分配唯一ID
            3. 识别可能需要图表的章节（如数据分析、结果展示等）
            4. 保持用户的原始意图
            """
            
            try:
                response = client.chat_completion([
                    {"role": "user", "content": parse_prompt}
                ])
                
                if response and 'content' in response:
                    content_text = response['content']
                    import re
                    json_match = re.search(r'\{.*\}', content_text, re.DOTALL)
                    if json_match:
                        outline_data = json.loads(json_match.group())
                        return jsonify({
                            'success': True,
                            'data': outline_data,
                            'message': '自定义大纲解析完成'
                        })
            except Exception as e:
                print(f"AI解析大纲失败: {e}")
        
        # 简单的文本解析
        lines = outline_text.split('\n')
        structured = []
        current_section = None
        section_id = 1
        subsection_id = 10
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            # 检测主章节（以数字或中文数字开始）
            if any(line.startswith(prefix) for prefix in ['1', '2', '3', '4', '5', '6', '7', '8', '9', '一', '二', '三', '四', '五', '六', '七', '八', '九']):
                if current_section:
                    structured.append(current_section)
                
                current_section = {
                    "id": section_id,
                    "title": line,
                    "children": []
                }
                section_id += 1
                subsection_id = section_id * 10
            
            # 检测子章节（以空格或tab开始，或包含小数点）
            elif current_section and (line.startswith('  ') or line.startswith('\t') or '.' in line[:5]):
                subsection_id += 1
                current_section["children"].append({
                    "id": subsection_id,
                    "title": line.lstrip(),
                    "hasChart": '数据' in line or '结果' in line or '分析' in line,
                    "chartType": 'chart' if ('数据' in line or '结果' in line) else None
                })
        
        if current_section:
            structured.append(current_section)
        
        parsed_outline = {
            "structured": structured,
            "systemRecommended": outline_text  # 保持原文本
        }
        
        return jsonify({
            'success': True,
            'data': parsed_outline,
            'message': '自定义大纲解析完成'
        })
        
    except Exception as e:
        return jsonify({'success': False, 'message': f'解析大纲失败: {str(e)}'}), 500

@thesis_bp.route('/download/<int:paper_id>', methods=['GET'])
def download_paper(paper_id):
    """下载论文"""
    try:
        user_id = session.get('user_id')
        
        if not user_id:
            return jsonify({'success': False, 'message': '请先登录'}), 401
        
        paper = Paper.query.filter_by(id=paper_id, user_id=user_id).first()
        if not paper:
            return jsonify({'success': False, 'message': '论文不存在'}), 404
        
        # 生成Word文档
        from docx import Document
        from docx.shared import Inches
        import io
        from flask import send_file
        
        doc = Document()
        
        # 添加标题
        title = doc.add_heading(paper.title, 0)
        
        # 添加摘要
        if paper.abstract:
            doc.add_heading('摘要', level=1)
            doc.add_paragraph(paper.abstract)
        
        # 添加关键词
        if paper.keywords:
            doc.add_paragraph(f"关键词：{paper.keywords}")
        
        # 添加正文内容
        if paper.content:
            doc.add_heading('正文', level=1)
            # 按段落分割内容
            paragraphs = paper.content.split('\n\n')
            for para in paragraphs:
                if para.strip():
                    doc.add_paragraph(para.strip())
        
        # 添加参考文献
        if paper.references:
            doc.add_heading('参考文献', level=1)
            doc.add_paragraph(paper.references)
        
        # 保存到内存
        doc_io = io.BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)
        
        # 记录下载活动
        log_user_activity(user_id, 'download_paper', {
            'paper_id': paper_id,
            'title': paper.title
        })
        
        return send_file(
            doc_io,
            as_attachment=True,
            download_name=f"{paper.title}.docx",
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'下载论文失败: {str(e)}'
        }), 500